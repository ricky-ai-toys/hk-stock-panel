"""
Excel export service — export CCASS data, trading data, shareholder data to Excel.
"""
import os
import logging
from datetime import datetime
from io import BytesIO
from typing import Optional

logger = logging.getLogger(__name__)


def export_stock_to_excel(stock_code: str) -> bytes:
    """Export all stock data to a single Excel file with multiple sheets."""
    from db import get_stock_conn
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, numbers

    code = stock_code.zfill(5)
    wb = openpyxl.Workbook()

    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    pct_fmt = '0.00%'
    num_fmt = '#,##0'
    price_fmt = '#,##0.00'

    def style_header(ws, row_num, cols):
        for col_idx in range(1, cols + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

    def autosize(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    val = str(cell.value) if cell.value is not None else ''
                    max_len = max(max_len, min(len(val), 40))
                except:
                    pass
            ws.column_dimensions[col_letter].width = max_len + 4

    with get_stock_conn(code) as conn:
        # Sheet 1: Overview / Stock Info
        ws1 = wb.active
        ws1.title = '基本信息'
        info_rows = conn.execute("SELECT key, value FROM stock_info").fetchall()
        ws1.append(['指标', '值'])
        style_header(ws1, 1, 2)
        for r in info_rows:
            ws1.append([r['key'], r['value']])
        autosize(ws1)

        # Sheet 2: Daily Trading
        ws2 = wb.create_sheet('行情数据')
        trading = conn.execute("""
            SELECT trade_date, open, high, low, close, volume, turnover, change_pct
            FROM daily_trading ORDER BY trade_date DESC
        """).fetchall()
        ws2.append(['日期', '开盘', '最高', '最低', '收盘', '成交量', '成交额(HKD)', '涨跌幅%'])
        style_header(ws2, 1, 8)
        for r in trading:
            pct = r['change_pct'] / 100 if r['change_pct'] is not None else None
            ws2.append([
                r['trade_date'], r['open'], r['high'], r['low'], r['close'],
                r['volume'], r['turnover'], pct
            ])
        for col in ['B', 'C', 'D', 'E']:
            for cell in ws2[col][1:]:
                cell.number_format = price_fmt
        for cell in ws2['F'][1:]:
            cell.number_format = num_fmt
        for cell in ws2['G'][1:]:
            cell.number_format = num_fmt
        for cell in ws2['H'][1:]:
            cell.number_format = '0.00%'
        autosize(ws2)

        # Sheet 3: CCASS Total + HK Connect
        ws3 = wb.create_sheet('CCASS总览')
        ccass = conn.execute("""
            SELECT ct.trade_date, ct.total_shares, ct.pct_of_issued,
                   hk.shares as hk_connect_shares, hk.pct_of_issued as hk_connect_pct,
                   hk.shares_change as hk_connect_change
            FROM ccass_total ct
            LEFT JOIN hk_connect hk ON ct.trade_date = hk.trade_date
            ORDER BY ct.trade_date DESC
        """).fetchall()
        ws3.append(['日期', 'CCASS总持仓(股)', '占总股本%', '南向持股(股)', '南向占比%', '南向变动(股)'])
        style_header(ws3, 1, 6)
        for r in ccass:
            ws3.append([
                r['trade_date'], r['total_shares'],
                r['pct_of_issued'] / 100 if r['pct_of_issued'] else None,
                r['hk_connect_shares'],
                r['hk_connect_pct'] / 100 if r['hk_connect_pct'] else None,
                r['hk_connect_change'],
            ])
        for col in ['B', 'D', 'F']:
            for cell in ws3[col][1:]:
                cell.number_format = num_fmt
        for col in ['C', 'E']:
            for cell in ws3[col][1:]:
                cell.number_format = '0.00%'
        autosize(ws3)

        # Sheet 4: CCASS Participants (latest date)
        ws4 = wb.create_sheet('席位持仓(最新)')
        latest_date = conn.execute("SELECT MAX(trade_date) as d FROM ccass_participants").fetchone()
        if latest_date and latest_date['d']:
            participants = conn.execute("""
                SELECT cp.participant_id, cp.participant_name, cp.shares, cp.pct_of_issued,
                       cdc.shares_change, cdc.shares_change_3d, cdc.shares_change_5d,
                       sa.shareholder_name
                FROM ccass_participants cp
                LEFT JOIN ccass_daily_changes cdc ON cp.participant_id = cdc.participant_id AND cp.trade_date = cdc.trade_date
                LEFT JOIN shareholder_attribution sa ON cp.participant_id = sa.participant_id
                WHERE cp.trade_date = ?
                ORDER BY cp.shares DESC
            """, (latest_date['d'],)).fetchall()
            ws4.append(['席位编号', '席位名称', '持仓(股)', '占比%', '日变动(股)', '3日变动(股)', '5日变动(股)', '归因股东'])
            style_header(ws4, 1, 8)
            for r in participants:
                ws4.append([
                    r['participant_id'], r['participant_name'], r['shares'],
                    r['pct_of_issued'] / 100 if r['pct_of_issued'] else None,
                    r['shares_change'], r['shares_change_3d'], r['shares_change_5d'],
                    r['shareholder_name'] or '',
                ])
            for col in ['C']:
                for cell in ws4[col][1:]:
                    cell.number_format = num_fmt
            for col in ['D']:
                for cell in ws4[col][1:]:
                    cell.number_format = '0.00%'
            for col in ['E', 'F', 'G']:
                for cell in ws4[col][1:]:
                    cell.number_format = num_fmt
        autosize(ws4)

        # Sheet 5: CCASS Daily Changes (latest date)
        ws5 = wb.create_sheet('席位日变动(最新)')
        if latest_date and latest_date['d']:
            changes = conn.execute("""
                SELECT cdc.participant_id, cdc.participant_name, cdc.shares,
                       cdc.shares_change, cdc.shares_change_3d, cdc.shares_change_5d,
                       sa.shareholder_name
                FROM ccass_daily_changes cdc
                LEFT JOIN shareholder_attribution sa ON cdc.participant_id = sa.participant_id
                WHERE cdc.trade_date = ?
                ORDER BY ABS(cdc.shares_change) DESC
            """, (latest_date['d'],)).fetchall()
            ws5.append(['席位编号', '席位名称', '持仓(股)', '日变动(股)', '3日变动(股)', '5日变动(股)', '归因股东'])
            style_header(ws5, 1, 7)
            for r in changes:
                ws5.append([
                    r['participant_id'], r['participant_name'], r['shares'],
                    r['shares_change'], r['shares_change_3d'], r['shares_change_5d'],
                    r['shareholder_name'] or '',
                ])
            for col in ['C', 'D', 'E', 'F']:
                for cell in ws5[col][1:]:
                    cell.number_format = num_fmt
        autosize(ws5)

        # Sheet 6: Shareholder Attribution
        ws6 = wb.create_sheet('股东归因')
        attrs = conn.execute("SELECT * FROM shareholder_attribution").fetchall()
        if attrs:
            ws6.append(['股东名称', '席位编号', '席位名称', '基线持股', '投资者类型', '成本价(HKD)', '解禁日', '置信度'])
            style_header(ws6, 1, 8)
            for r in attrs:
                ws6.append([
                    r['shareholder_name'], r['participant_id'], r['participant_name'],
                    r['shares_baseline'], r['investor_type'], r['cost_hkd'],
                    r['unlock_date'], r['confidence'],
                ])
            for col in ['D']:
                for cell in ws6[col][1:]:
                    cell.number_format = num_fmt
            for col in ['F']:
                for cell in ws6[col][1:]:
                    cell.number_format = price_fmt
        autosize(ws6)

        # Sheet 7: Shareholder Selloff
        ws7 = wb.create_sheet('股东减持追踪')
        selloffs = conn.execute("""
            SELECT trade_date, shareholder_name, shares_change, shares_remaining, estimated_proceeds
            FROM shareholder_selloff ORDER BY trade_date DESC, ABS(shares_change) DESC
        """).fetchall()
        if selloffs:
            ws7.append(['日期', '股东名称', '日变动(股)', '剩余持仓(股)', '估算金额(HKD)'])
            style_header(ws7, 1, 5)
            for r in selloffs:
                ws7.append([r['trade_date'], r['shareholder_name'], r['shares_change'], r['shares_remaining'], r['estimated_proceeds']])
            for col in ['C', 'D', 'E']:
                for cell in ws7[col][1:]:
                    cell.number_format = num_fmt
        autosize(ws7)

        # Sheet 8: Unlock Calendar
        ws8 = wb.create_sheet('解禁日历')
        unlocks = conn.execute("SELECT * FROM unlock_calendar ORDER BY unlock_date").fetchall()
        if unlocks:
            ws8.append(['解禁日期', '股东名称', '持股(股)', '成本价(HKD)', '券商', '类型', '状态'])
            style_header(ws8, 1, 7)
            for r in unlocks:
                ws8.append([
                    r['unlock_date'], r['shareholder_name'], r['shares'],
                    r['cost_hkd'], r['broker'], r['investor_type'], r['unlock_status'],
                ])
            for col in ['C']:
                for cell in ws8[col][1:]:
                    cell.number_format = num_fmt
        autosize(ws8)

        # Sheet 9: Announcements (if exists)
        try:
            anns = conn.execute("SELECT announce_date, title, category FROM announcements ORDER BY announce_date DESC").fetchall()
            if anns:
                ws9 = wb.create_sheet('公司公告')
                ws9.append(['日期', '标题', '类型'])
                style_header(ws9, 1, 3)
                for r in anns:
                    ws9.append([r['announce_date'], r['title'], r['category']])
                autosize(ws9)
        except Exception:
            pass

    # Save to bytes
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


if __name__ == '__main__':
    import sys
    logging.basicConfig(level=logging.INFO)
    code = sys.argv[1] if len(sys.argv) > 1 else '00000'
    data = export_stock_to_excel(code)
    outpath = f'/tmp/{code}_export.xlsx'
    with open(outpath, 'wb') as f:
        f.write(data)
    print(f"Exported to {outpath} ({len(data)} bytes)")
